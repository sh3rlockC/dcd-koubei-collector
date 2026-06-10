import importlib.util
import sys
from pathlib import Path

from openpyxl import load_workbook


SCRIPT_PATH = Path(__file__).resolve().parents[1] / "scripts" / "export_dcd_koubei.py"


def load_module():
    spec = importlib.util.spec_from_file_location("export_dcd_koubei", SCRIPT_PATH)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


def test_filter_known_rows_skips_known_links_and_keeps_missing_links():
    module = load_module()
    rows = [
        {"用户名": "历史车主", "来源链接": "https://www.dongchedi.com/koubei/1", "评价全文": "old"},
        {"用户名": "新车主", "来源链接": "https://www.dongchedi.com/koubei/2", "评价全文": "new"},
        {"用户名": "无链接车主", "来源链接": "", "评价全文": "fallback hash later"},
    ]

    filtered, known_count = module.filter_known_rows(rows, {"https://www.dongchedi.com/koubei/1"})

    assert known_count == 1
    assert [row["用户名"] for row in filtered] == ["新车主", "无链接车主"]


def test_main_stops_after_consecutive_known_pages(monkeypatch, tmp_path: Path):
    module = load_module()
    output = tmp_path / "DCD口碑_测试车.xlsx"
    known_file = tmp_path / "known.txt"
    known_file.write_text("https://www.dongchedi.com/koubei/1\n", encoding="utf-8")
    calls = []

    def fake_detect_total_pages(series_id):
        return 5, "测试车"

    def fake_collect_page(series_id, page):
        calls.append(page)
        if page == 3:
            return (
                [
                    {
                        "用户名": "新车主",
                        "评价车型": "测试车",
                        "发布时间": "2026-05-03",
                        "评价全文": "第三页新增，不应触达",
                        "来源链接": "https://www.dongchedi.com/koubei/3",
                        "抓取页码": "3",
                    }
                ],
                [],
                {"page": page, "series_name": "测试车", "page_count": 1},
                0,
            )
        return (
            [
                {
                    "用户名": f"历史车主{page}",
                    "评价车型": "测试车",
                    "发布时间": "2026-05-01",
                    "评价全文": "历史评论",
                    "来源链接": "https://www.dongchedi.com/koubei/1",
                    "抓取页码": str(page),
                }
            ],
            [],
            {"page": page, "series_name": "测试车", "page_count": 1},
            0,
        )

    monkeypatch.setattr(module, "detect_total_pages", fake_detect_total_pages)
    monkeypatch.setattr(module, "collect_page", fake_collect_page)
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "export_dcd_koubei.py",
            "--series-id",
            "25398",
            "--output",
            str(output),
            "--known-links-file",
            str(known_file),
            "--max-scan-pages",
            "10",
            "--stop-after-known-pages",
            "2",
            "--quiet",
        ],
    )

    module.main()

    assert calls == [1, 2]
    workbook = load_workbook(output)
    assert workbook.active.max_row == 1
    validation = module.json.loads(output.with_suffix(".validation.json").read_text(encoding="utf-8"))
    assert validation["incremental"]["pages_scanned"] == 2
    assert validation["incremental"]["stop_reason"] == "known_pages"
