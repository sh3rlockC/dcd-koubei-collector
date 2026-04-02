#!/usr/bin/env python3
"""导出懂车帝车型口碑为 Excel。

当前实现：
- 直接请求懂车帝车型口碑分页 URL
- 从页面内嵌 __NEXT_DATA__ 提取结构化 review_list
- 支持自动探测总页数
- 支持默认文件名：DCD口碑_车型_日期.xlsx
- 导出 Excel + validation.json

适用前提：
- 页面 HTML 中仍包含可直接解析的 __NEXT_DATA__
- 分页 URL 形如：
  https://www.dongchedi.com/auto/series/score/{series_id}-x-S0-x-x-x-{page}
"""

import argparse
import datetime as dt
import json
import re
import sys
import time
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple, Union


class ProgressTracker:
    def __init__(self, label: str, progress_path: Optional[Path] = None, quiet: bool = False):
        self.label = label
        self.progress_path = progress_path
        self.quiet = quiet
        self.state: Dict[str, Any] = {
            'label': label,
            'stage': 'init',
            'current': 0,
            'total': 0,
            'percent': 0,
            'success': 0,
            'retry': 0,
            'failed': 0,
            'records': 0,
            'message': '',
            'updated_at': '',
            'overall': {'current': 0, 'total': 0, 'percent': 0},
            'current_page': 0,
            'page_range': {'start': 0, 'end': 0},
            'output_path': '',
            'validation_path': '',
            'failed_pages': [],
        }

    def _render_bar(self, current: int, total: int, width: int = 24) -> str:
        if total <= 0:
            return '[' + '.' * width + ']'
        filled = min(width, int(width * current / total))
        return '[' + '#' * filled + '.' * (width - filled) + ']'

    def emit(self):
        self.state['updated_at'] = dt.datetime.now().isoformat(timespec='seconds')
        if self.progress_path:
            self.progress_path.write_text(json.dumps(self.state, ensure_ascii=False, indent=2), encoding='utf-8')
        if self.quiet:
            return
        overall = self.state.get('overall') or {}
        bar = self._render_bar(int(overall.get('current', 0)), int(overall.get('total', 0)))
        total = overall.get('total') or '?'
        page_range = self.state.get('page_range') or {}
        print(
            f"{self.state['stage']} {bar} 总体 {overall.get('current', 0)}/{total} ({overall.get('percent', 0)}%) | "
            f"页码 {self.state.get('current_page', 0)}/{page_range.get('end', 0)} | ok {self.state['success']} retry {self.state['retry']} "
            f"fail {self.state['failed']} rows {self.state['records']} | {self.state['message']}",
            flush=True,
        )

    def update(self, **kwargs):
        self.state.update(kwargs)
        total = int(self.state.get('total') or 0)
        current = int(self.state.get('current') or 0)
        percent = int(current * 100 / total) if total > 0 else 0
        self.state['percent'] = percent
        self.state['overall'] = {'current': current, 'total': total, 'percent': percent}
        self.emit()

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

HEADERS = [
    '用户名', '用户标签', '评价车型', '懂车分', '发布时间', '用户评分', '续航', '购车时间',
    '裸车价', '购车地', '评价全文', '来源链接', '抓取页码'
]

UA = 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/123.0.0.0 Safari/537.36'


def fetch_html(url: str, retries: int = 2, timeout: int = 20) -> Tuple[str, int]:
    last_err = None
    retry_count = 0
    for i in range(retries + 1):
        try:
            r = requests.get(url, headers={'User-Agent': UA}, timeout=timeout)
            r.raise_for_status()
            return r.text, retry_count
        except Exception as e:
            last_err = e
            if i < retries:
                retry_count += 1
                time.sleep(1)
    raise RuntimeError(str(last_err))


def extract_next_data(html: str) -> Dict[str, Any]:
    m = re.search(r'__NEXT_DATA__[^>]*>(\{.*?\})</script>', html, re.S)
    if not m:
        raise RuntimeError('页面中未找到 __NEXT_DATA__')
    return json.loads(m.group(1))


def series_page_url(series_id: int, page: int) -> str:
    return f'https://www.dongchedi.com/auto/series/score/{series_id}-x-S0-x-x-x-{page}'


def infer_series_id_from_url(url: str) -> int:
    m = re.search(r'/auto/series/score/(\d+)-', url)
    if not m:
        raise ValueError(f'无法从 URL 提取 series_id: {url}')
    return int(m.group(1))


def normalize_score(raw: Any) -> str:
    if raw in (None, ''):
        return ''
    try:
        v = float(raw)
    except Exception:
        return str(raw)
    if v >= 100:
        return f'{v / 100:.2f}'.rstrip('0').rstrip('.')
    return f'{v:.2f}'.rstrip('0').rstrip('.')


def ts_to_date(ts: Any) -> str:
    if not ts:
        return ''
    try:
        return dt.datetime.fromtimestamp(int(ts)).strftime('%Y-%m-%d')
    except Exception:
        return ''


def review_url(gid: Union[str, int]) -> str:
    return f'https://www.dongchedi.com/koubei/{gid}'


def safe_filename(name: str) -> str:
    name = re.sub(r'[\\/:*?"<>|]+', '_', name)
    name = re.sub(r'\s+', '_', name).strip('_')
    return name or '未命名车型'


def default_output_name(series_name: str) -> str:
    today = dt.datetime.now().strftime('%Y-%m-%d')
    return f'DCD口碑_{safe_filename(series_name)}_{today}.xlsx'


def build_row(item: Dict[str, Any], page: int) -> Dict[str, str]:
    buy = item.get('buy_car_info') or {}
    score = item.get('score_info') or {}
    user = item.get('user_info') or {}

    row = {
        '用户名': user.get('name', '') or '',
        '用户标签': user.get('tag', '') or '',
        '评价车型': ' '.join(x for x in [item.get('series_name', ''), item.get('car_name', '')] if x).strip(),
        '懂车分': normalize_score(score.get('score')),
        '发布时间': ts_to_date(item.get('create_time')),
        '用户评分': normalize_score(score.get('score')),
        '续航': buy.get('continuation', '') or '',
        '购车时间': buy.get('bought_time', '') or '',
        '裸车价': buy.get('price', '') or '',
        '购车地': buy.get('location', '') or '',
        '评价全文': (item.get('content') or '').strip(),
        '来源链接': review_url(item.get('gid_str') or item.get('gid') or ''),
        '抓取页码': str(page),
    }
    return row


def is_valid_row(row: Dict[str, str]) -> Tuple[bool, List[str]]:
    missing = [k for k in ['评价车型', '发布时间', '评价全文', '来源链接'] if not row.get(k)]
    return len(missing) == 0, missing


def collect_page(series_id: int, page: int) -> Tuple[List[Dict[str, str]], List[Dict[str, Any]], Dict[str, Any], int]:
    url = series_page_url(series_id, page)
    html, retry_count = fetch_html(url)
    data = extract_next_data(html)
    pp = data['props']['pageProps']
    review_data = pp.get('reviewListData') or {}
    review_list = review_data.get('review_list') or []

    rows: List[Dict[str, str]] = []
    anomalies: List[Dict[str, Any]] = []

    for item in review_list:
        row = build_row(item, page)
        ok, missing = is_valid_row(row)
        if ok:
            rows.append(row)
        else:
            anomalies.append({
                'type': 'missing_required_fields',
                'page': page,
                'missing': missing,
                'gid': item.get('gid_str') or item.get('gid'),
                'preview': row,
            })

    meta = {
        'page': page,
        'url': url,
        'series_name': pp.get('seriesName', ''),
        'page_count': len(review_list),
        'has_more': bool(review_data.get('has_more')),
        'total_count': review_data.get('total_count'),
    }
    return rows, anomalies, meta, retry_count


def detect_total_pages(series_id: int, max_scan: int = 200) -> Tuple[int, Optional[str]]:
    last_non_empty = 0
    detected_series_name = None

    for page in range(1, max_scan + 1):
        try:
            html, _ = fetch_html(series_page_url(series_id, page))
            data = extract_next_data(html)
            pp = data['props']['pageProps']
            review_data = pp.get('reviewListData') or {}
            review_list = review_data.get('review_list') or []
            has_more = bool(review_data.get('has_more'))

            if pp.get('seriesName') and not detected_series_name:
                detected_series_name = pp.get('seriesName')

            if review_list:
                last_non_empty = page
            else:
                break

            if not has_more:
                return page, detected_series_name
        except Exception:
            break

    return max(last_non_empty, 1), detected_series_name


def apply_sheet_style(ws):
    for c in ws[1]:
        c.font = Font(bold=True)

    for col in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'M']:
        ws.column_dimensions[col].width = 16
    ws.column_dimensions['K'].width = 100
    ws.column_dimensions['L'].width = 60

    for row in ws.iter_rows(min_row=2):
        row[10].alignment = Alignment(wrap_text=True, vertical='top')
        row[11].alignment = Alignment(wrap_text=True, vertical='top')


def write_xlsx(out_path: Path, rows: List[Dict[str, str]]):
    wb = Workbook()
    ws = wb.active
    ws.title = '口碑明细'
    ws.append(HEADERS)

    for row in rows:
        ws.append([row.get(h, '') for h in HEADERS])

    apply_sheet_style(ws)
    wb.save(out_path)


def read_existing_rows(path: Path) -> List[Dict[str, str]]:
    wb = load_workbook(path)
    ws = wb['口碑明细'] if '口碑明细' in wb.sheetnames else wb.active
    header = [cell.value for cell in ws[1]]
    rows: List[Dict[str, str]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        item = {}
        for i, key in enumerate(header):
            if key:
                item[str(key)] = '' if row[i] is None else str(row[i])
        if any(item.values()):
            rows.append(item)
    return rows


def merge_rows(existing_rows: List[Dict[str, str]], new_rows: List[Dict[str, str]], mode: str = 'keep-extra') -> List[Dict[str, str]]:
    if mode == 'strict':
        by_link: Dict[str, Dict[str, str]] = {}
        order: List[str] = []
        for row in new_rows:
            key = row.get('来源链接') or ''
            if not key:
                continue
            if key not in by_link:
                order.append(key)
            by_link[key] = row
        return [by_link[key] for key in order]

    merged: Dict[str, Dict[str, str]] = {}
    order: List[str] = []

    for row in existing_rows:
        key = row.get('来源链接') or ''
        if not key:
            continue
        if key not in merged:
            order.append(key)
        merged[key] = row

    for row in new_rows:
        key = row.get('来源链接') or ''
        if not key:
            continue
        if key not in merged:
            order.append(key)
        merged[key] = row

    return [merged[key] for key in order]


def write_validation_report(path: Path, report: Dict[str, Any]):
    path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')


def load_retry_pages(path: Path) -> List[int]:
    data = json.loads(path.read_text(encoding='utf-8'))
    pages: List[int] = []
    for item in data:
        page = item.get('page')
        if isinstance(page, int) and page > 0:
            pages.append(page)
        elif isinstance(page, str) and page.isdigit() and int(page) > 0:
            pages.append(int(page))
    return sorted(set(pages))


def main():
    parser = argparse.ArgumentParser(description='Export DCD koubei to xlsx')
    parser.add_argument('--url', help='懂车帝口碑页 URL，可自动提取 series_id')
    parser.add_argument('--series-id', type=int, help='车型 series_id')
    parser.add_argument('--start-page', type=int, default=1)
    parser.add_argument('--end-page', type=int, help='结束页；不传时自动探测')
    parser.add_argument('--output', help='输出 xlsx 路径；不传时默认用 DCD口碑_车型_日期.xlsx')
    parser.add_argument('--progress-file', help='进度 JSON 输出路径；默认与 output 同目录同名 .progress.json')
    parser.add_argument('--quiet', action='store_true', help='静默模式，仅写 progress json，不打印进度行')
    parser.add_argument('--retry-failed-pages', help='从 *.failed-pages.json 读取失败页，仅补抓这些页')
    parser.add_argument('--merge-into', help='将补抓结果合并进已有 Excel，按 来源链接 覆盖旧记录')
    parser.add_argument('--merge-mode', choices=['keep-extra', 'strict'], default='keep-extra', help='merge-into 时的合并模式：keep-extra 保留旧表其他记录；strict 仅保留本轮新结果')
    args = parser.parse_args()

    if not args.series_id and not args.url:
        raise SystemExit('必须提供 --series-id 或 --url')

    series_id = args.series_id or infer_series_id_from_url(args.url)

    if args.start_page < 1:
        raise SystemExit('invalid page range: require start-page >= 1')

    detected_series_name = None
    auto_detected = False
    retry_pages: List[int] = []
    if args.retry_failed_pages:
        retry_pages = load_retry_pages(Path(args.retry_failed_pages).resolve())
        if not retry_pages:
            raise SystemExit('retry-failed-pages 中没有可用页码')
        end_page = max(retry_pages)
        args.start_page = min(retry_pages)
    else:
        end_page = args.end_page
        if end_page is None:
            end_page, detected_series_name = detect_total_pages(series_id)
            auto_detected = True

    if end_page < args.start_page:
        raise SystemExit('invalid page range: require end-page >= start-page')

    all_rows: List[Dict[str, str]] = []
    anomalies: List[Dict[str, Any]] = []
    page_counts: Dict[str, int] = {}
    page_meta: List[Dict[str, Any]] = []
    series_name = detected_series_name or ''

    output = args.output or default_output_name(series_name or str(series_id))
    out_path = Path(output).resolve()
    progress_path = Path(args.progress_file).resolve() if args.progress_file else out_path.with_suffix('.progress.json')
    tracker = ProgressTracker(label=f'dcd:{series_id}', progress_path=progress_path, quiet=args.quiet)
    target_pages = retry_pages or list(range(args.start_page, end_page + 1))
    total_steps = len(target_pages) + 2
    failed_page_list: List[Dict[str, Any]] = []
    tracker.update(
        stage='初始化',
        current=0,
        total=total_steps,
        current_page=0,
        page_range={'start': min(target_pages), 'end': max(target_pages)},
        output_path=str(out_path),
        failed_pages=failed_page_list,
        message='准备开始抓取懂车帝口碑' + ('（失败页补抓）' if retry_pages else ''),
    )

    success_pages = 0
    retry_total = 0
    failed_pages = 0
    for index, page in enumerate(target_pages, start=1):
        current_step = index
        tracker.update(
            stage='抓取页面',
            current=current_step,
            total=total_steps,
            success=success_pages,
            retry=retry_total,
            failed=failed_pages,
            records=len(all_rows),
            current_page=page,
            page_range={'start': min(target_pages), 'end': max(target_pages)},
            failed_pages=failed_page_list,
            message=f'第 {page} 页',
        )
        try:
            rows, bad, meta, retry_count = collect_page(series_id, page)
            all_rows.extend(rows)
            anomalies.extend(bad)
            page_counts[str(page)] = len(rows)
            page_meta.append(meta)
            retry_total += retry_count
            success_pages += 1
            if not series_name:
                series_name = meta.get('series_name', '')
            tracker.update(
                stage='抓取页面',
                current=current_step,
                total=total_steps,
                success=success_pages,
                retry=retry_total,
                failed=failed_pages,
                records=len(all_rows),
                current_page=page,
                page_range={'start': min(target_pages), 'end': max(target_pages)},
                failed_pages=failed_page_list,
                message=f'第 {page} 页完成，新增 {len(rows)} 条',
            )
        except Exception as e:
            failure = {
                'type': 'page_fetch_failed',
                'page': page,
                'reason': str(e),
                'url': series_page_url(series_id, page),
            }
            anomalies.append(failure)
            failed_page_list.append(failure)
            page_counts[str(page)] = 0
            failed_pages += 1
            tracker.update(
                stage='抓取页面',
                current=current_step,
                total=total_steps,
                success=success_pages,
                retry=retry_total,
                failed=failed_pages,
                records=len(all_rows),
                current_page=page,
                page_range={'start': min(target_pages), 'end': max(target_pages)},
                failed_pages=failed_page_list,
                message=f'第 {page} 页失败',
            )
    final_rows = all_rows
    merge_summary = None
    if args.merge_into:
        merge_path = Path(args.merge_into).resolve()
        existing_rows = read_existing_rows(merge_path)
        final_rows = merge_rows(existing_rows, all_rows, mode=args.merge_mode)
        merge_summary = {
            'merge_into': str(merge_path),
            'merge_mode': args.merge_mode,
            'existing_rows': len(existing_rows),
            'new_rows': len(all_rows),
            'merged_rows': len(final_rows),
        }

    tracker.update(
        stage='导出结果',
        current=total_steps - 1,
        total=total_steps,
        success=success_pages,
        retry=retry_total,
        failed=failed_pages,
        records=len(final_rows),
        current_page=target_pages[-1],
        page_range={'start': min(target_pages), 'end': max(target_pages)},
        failed_pages=failed_page_list,
        message='写入 Excel 与 validation.json',
    )
    write_xlsx(out_path, final_rows)

    report = {
        'ok': len([a for a in anomalies if a.get('type') == 'page_fetch_failed']) == 0,
        'series_id': series_id,
        'series_name': series_name,
        'input': {
            'url': args.url,
            'start_page': args.start_page,
            'end_page': end_page,
            'end_page_auto_detected': auto_detected,
            'retry_failed_pages': args.retry_failed_pages,
            'target_pages': target_pages,
        },
        'total_rows': len(final_rows),
        'page_counts': page_counts,
        'page_meta': page_meta,
        'anomaly_count': len(anomalies),
        'anomalies': anomalies,
        'merge': merge_summary,
    }
    validation_path = out_path.with_suffix('.validation.json')
    write_validation_report(validation_path, report)
    failed_pages_path = out_path.with_suffix('.failed-pages.json')
    failed_pages_path.write_text(json.dumps(failed_page_list, ensure_ascii=False, indent=2), encoding='utf-8')

    tracker.update(
        stage='完成',
        current=total_steps,
        total=total_steps,
        success=success_pages,
        retry=retry_total,
        failed=failed_pages,
        records=len(final_rows),
        current_page=target_pages[-1],
        page_range={'start': min(target_pages), 'end': max(target_pages)},
        validation_path=str(validation_path),
        failed_pages=failed_page_list,
        message=f'导出完成: {out_path.name}',
    )

    print(f'Wrote workbook: {out_path}')
    print(f'Validation: {validation_path}')
    print(f'Series: {series_name or series_id}')
    if retry_pages:
        print(f'Retry pages: {target_pages}')
    else:
        print(f'Page range: {args.start_page}-{end_page}' + (' (auto-detected)' if auto_detected else ''))
    print(f'Total rows: {len(final_rows)}')
    if merge_summary:
        print(f"Merged into existing workbook: {merge_summary['merge_into']}")
        print(f"Existing rows: {merge_summary['existing_rows']}, New rows: {merge_summary['new_rows']}, Final rows: {merge_summary['merged_rows']}")
    print(f'Anomalies: {len(anomalies)}')


if __name__ == '__main__':
    try:
        main()
    except KeyboardInterrupt:
        sys.exit(130)
